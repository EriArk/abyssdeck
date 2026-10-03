"""Independent OOXML fixture and round-trip validation; run with openpyxl 3.1.5."""
import sys
from pathlib import Path
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.chart import BarChart, Reference
from openpyxl.worksheet.table import Table
from openpyxl.utils.escape import unescape

target = Path(sys.argv[2])
if sys.argv[1] == 'create':
    book = Workbook()
    sheet = book.active
    sheet.title = 'Budget'
    sheet.append(['Name', 'Amount', 'Total', 'Date', 'Note', 'Flag'])
    sheet.append(['Alpha', 12.5, '=B2*2', 45200, 'untouched', True])
    sheet.append(['Hidden row', 7, '=B3*2', 45201, 'keep me', False])
    for n in range(4, 64):
        sheet.append([f'Row {n}', n])
    sheet['B2'].font = Font(bold=True, color='112233')
    sheet['B2'].fill = PatternFill('solid', fgColor='FFFF00')
    sheet['D2'].number_format = 'yyyy-mm-dd'
    sheet.merge_cells('G2:H2')
    sheet['G2'] = 'Merged'
    sheet.freeze_panes = 'B2'
    sheet.row_dimensions[3].hidden = True
    sheet['E2'].hyperlink = 'https://example.com/'
    chart = BarChart()
    chart.add_data(Reference(sheet, min_col=2, min_row=1, max_row=4), titles_from_data=True)
    sheet.add_chart(chart, 'J2')
    other = book.create_sheet('Other sheet')
    other.append(['Independent', '=Budget!B2+1'])
    other['A1'].font = Font(italic=True)
    protected = book.create_sheet('Protected')
    protected['A1'] = 'Read only'
    protected.protection.sheet = True
    table = book.create_sheet('Structured')
    table.append(['Key', 'Value'])
    table.append(['Preserve', 123])
    table.add_table(Table(displayName='DataTable', ref='A1:B2'))
    book.save(target)
else:
    book = load_workbook(target)
    sheet = book['Budget']
    # openpyxl's inlineStr parser leaves SpreadsheetML escapes encoded; its standard
    # escape decoder provides the logical string (one pass, preserving literal _x000A_).
    assert unescape(sheet['A2'].value) == 'Лазарь 🙂 _x000A_\r\nline'
    assert sheet['B2'].value == 123.75
    assert sheet['B2'].font.bold and sheet['B2'].fill.fgColor.rgb == '00FFFF00'
    assert sheet['D2'].number_format == 'yyyy-mm-dd'
    assert sheet['C2'].value == '=B2*2'
    assert sheet['E2'].hyperlink.target == 'https://example.com/'
    assert sheet['A3'].value == 'Hidden row' and sheet.row_dimensions[3].hidden
    assert sheet.freeze_panes == 'B2' and sheet._charts
    assert 'G2:H2' in sheet.merged_cells
    assert book['Other sheet']['B1'].value == '=Budget!B2+1'
    assert book['Other sheet']['A1'].font.italic
    assert book['Protected'].protection.sheet
    assert 'DataTable' in book['Structured'].tables
    assert book.calculation.fullCalcOnLoad and book.calculation.forceFullCalc
    print('Independent openpyxl reopen: values, Unicode/CRLF, formulas, styles, dates, links, charts, tables, protection and other sheets passed')
